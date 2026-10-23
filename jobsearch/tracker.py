"""The Excel tracker, output/Job Tracker.xlsx.

Sheets
  Today's applications  rebuilt by every build from that day's picks. The Status, Date applied
                        and Notes typed here survive a same-day rebuild and move with the row.
  All applications      every job ever shortlisted. Rows are only appended, never rewritten,
                        so edits, extra columns and formatting made in Excel are kept.
  Summary               counts by status and follow-ups due (formulas, rebuilt each time).
Each row carries the job key in the last column, which is how rows are matched.
"""
import datetime as dt
import os
import shutil
from urllib.parse import quote, unquote

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from . import config, paths
from .textutil import detect_region, fingerprint

TODAY, ALL, SUMMARY = "Today's applications", "All applications", "Summary"
STATUSES = ["To apply", "Applied", "Interview", "Offer", "Rejected", "No response", "Withdrawn", "Skipped"]
STATUS_FILL = {"Applied": "DDEBF7", "Interview": "FFE699", "Offer": "A9D08E", "Rejected": "F4B183",
               "No response": "E7E6E6", "Withdrawn": "E7E6E6", "Skipped": "E7E6E6"}
TODAY_COLS = [("Date found", 11), ("Fit", 5), ("Job title", 34), ("Company", 24), ("Direct / Agency", 18),
              ("Location", 20), ("Advert", 11), ("Employer site", 12), ("Why apply", 70), ("CV to send", 26),
              ("Expected salary", 30), ("Meets target?", 14), ("Visa / right to work", 36), ("Cover letter", 12),
              ("Deadline", 11), ("Status", 12), ("Date applied", 12), ("Notes", 30), ("Key", 24)]
ALL_COLS = [("Date found", 11), ("Fit", 5), ("Job title", 34), ("Company", 24), ("Direct / Agency", 18),
            ("Location", 20), ("Advert", 11), ("Employer site", 12), ("CV sent", 26), ("Cover letter", 12),
            ("Deadline", 11), ("Status", 12), ("Date applied", 12), ("Follow up by", 12), ("Notes", 30),
            ("Why apply", 50), ("Key", 24)]
USER_FIELDS = ("Status", "Date applied", "Notes")  # typed by the candidate, carried over
ATS_SOURCES = {"greenhouse", "lever", "ashby", "smartrecruiters", "workable", "workday"}
FOLLOW_UP_DAYS = 7
MAX_ROWS = 5000
HEAD_FILL = PatternFill("solid", fgColor="1F3864")
LINK_FONT = Font(color="0563C1", underline="single")
DATE_FMT = "yyyy-mm-dd"
KEEP_BACKUPS = 20


class TrackerLocked(Exception):
    pass


# ---------------------------------------------------------------- reading
def _iso(v):
    if isinstance(v, (dt.datetime, dt.date)):
        return v.strftime("%Y-%m-%d")
    return str(v).strip()[:10] if v else ""


def _headers(ws):
    return {c.value: c.column for c in ws[1] if c.value}


def rows(ws):
    """Each data row as {header: value}, link columns as {header: (text, target)}."""
    heads = _headers(ws)
    out = []
    for r in range(2, ws.max_row + 1):
        row = {}
        for h, col in heads.items():
            cell = ws.cell(r, col)
            row[h] = (cell.value, cell.hyperlink.target) if cell.hyperlink else cell.value
        if any(v not in (None, "") for v in row.values()):
            row["_row"] = r
            out.append(row)
    return out


def value(v):
    return v[0] if isinstance(v, tuple) else v


def listed_before(date):
    """{job key or fingerprint: date found} for every job already in the tracker, except `date`'s."""
    if not paths.TRACKER_FILE.exists():
        return {}
    wb = load_workbook(paths.TRACKER_FILE)
    found = {}
    for name in (TODAY, ALL):
        if name not in wb.sheetnames:
            continue
        for row in rows(wb[name]):
            when = _iso(value(row.get("Date found")))
            if when == date:
                continue
            title, company = value(row.get("Job title")) or "", value(row.get("Company")) or ""
            region = detect_region(str(value(row.get("Location")) or ""))
            if row.get("Key"):
                found[row["Key"]] = when
            if title and company and region:
                found[fingerprint(company, title, region)] = when
    return found


# ---------------------------------------------------------------- writing helpers
def _link(cell, target, text):
    if target:
        cell.value, cell.hyperlink, cell.font = text, target, LINK_FONT


def _file_link(relative_to_output):
    """Relative link from the workbook's folder, so it survives moving the whole project."""
    return quote(relative_to_output.replace("\\", "/"), safe="/.-_()'")


def _new_sheet(wb, name, cols, index=None):
    ws = wb.create_sheet(name, index)
    ws.append([h for h, _ in cols])
    for i, (_, width) in enumerate(cols, 1):
        c = ws.cell(1, i)
        c.font, c.fill = Font(bold=True, color="FFFFFF"), HEAD_FILL
        c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}1"
    _status_rules(ws)
    return ws


def _status_rules(ws):
    """Status drop-down and colours for the whole column (added once per sheet)."""
    col = get_column_letter(_headers(ws)["Status"])
    rng = f"{col}2:{col}{MAX_ROWS}"
    formula = '"' + ",".join(STATUSES) + '"'
    if any(dv.formula1 == formula for dv in ws.data_validations.dataValidation):
        return
    dv = DataValidation(type="list", formula1=formula, allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(rng)
    for status, color in STATUS_FILL.items():
        ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=[f'"{status}"'],
                                                      fill=PatternFill("solid", fgColor=color)))


def _put(ws, values, links=(), dates=(), wrap=False):
    """Append a row given as {header: value}; links = {header: (target, text)}."""
    heads = _headers(ws)
    r = ws.max_row + 1
    for h, v in values.items():
        if h in heads and v not in (None, ""):
            ws.cell(r, heads[h], v)
    for h, (target, text) in dict(links).items():
        if h in heads:
            _link(ws.cell(r, heads[h]), target, text)
    for h in dates:
        if h in heads:
            ws.cell(r, heads[h]).number_format = DATE_FMT
    for c in ws[r]:
        c.alignment = Alignment(wrap_text=wrap, vertical="top")
    return r


def as_date(v):
    """Excel date / ISO text -> date; other text is returned unchanged."""
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    try:
        return dt.date.fromisoformat(str(v)[:10]) if v else None
    except ValueError:
        return v


def _cv_link(cv_id):
    path = config.cv_path(cv_id)
    try:
        rel = os.path.relpath(path, paths.OUTPUT_DIR)
    except ValueError:  # CV folder on another drive
        rel = str(path)
    return _file_link(rel) if not os.path.isabs(rel) else path.as_uri()


def _employer_link(pick):
    if pick.get("direct_link"):
        return pick["direct_link"]
    if pick["source"] in ATS_SOURCES:
        return ""  # the advert already is the employer's own page
    return next((a["link"] for a in pick.get("also_on", []) if a["source"] in ATS_SOURCES), "")


def _append_all(ws, row):
    """Move one 'Today' row (as read by rows()) to the bottom of All applications."""
    heads = _headers(ws)
    links = {h: (row[h][1], row[h][0]) for h in ("Advert", "Employer site", "Cover letter")
             if isinstance(row.get(h), tuple)}
    cv = row.get("CV to send")
    if isinstance(cv, tuple):
        links["CV sent"] = (cv[1], cv[0])
    values = {h: value(row.get(h)) for h in heads if h not in links and h != "Follow up by"}
    values["CV sent"] = values.get("CV sent") or value(cv)
    values["Status"] = values.get("Status") or "To apply"
    values["Date found"] = as_date(values.get("Date found"))
    values["Date applied"] = as_date(values.get("Date applied"))
    r = _put(ws, values, links, dates=("Date found", "Date applied", "Follow up by"))
    applied = get_column_letter(heads["Date applied"])
    ws.cell(r, heads["Follow up by"], f'=IF({applied}{r}="","",{applied}{r}+{FOLLOW_UP_DAYS})')


def _fix_moved_letters(ws):
    """Point links of letters that the weekly rotation moved at their new place."""
    col = _headers(ws).get("Cover letter")
    if not col:
        return
    for r in range(2, ws.max_row + 1):
        cell = ws.cell(r, col)
        target = cell.hyperlink.target if cell.hyperlink else ""
        if not target or "://" in target or (paths.OUTPUT_DIR / unquote(target)).exists():
            continue
        parts = unquote(target).split("/")  # Cover Letters/<date>/<file>
        if len(parts) == 3 and parts[0] == paths.LETTERS_DIR.name:
            new = paths.PREVIOUS_LETTERS_DIR / f"{parts[1]} - {parts[2]}"
            if new.exists():
                _link(cell, _file_link(new.relative_to(paths.OUTPUT_DIR).as_posix()), cell.value)


def _summary(wb):
    if SUMMARY in wb.sheetnames:
        wb.remove(wb[SUMMARY])
    ws = wb.create_sheet(SUMMARY)
    t_col = get_column_letter(_headers(wb[TODAY])["Status"])
    a_heads = _headers(wb[ALL])
    a_col, f_col = get_column_letter(a_heads["Status"]), get_column_letter(a_heads["Follow up by"])
    today_ref, all_ref = "'Today''s applications'", "'All applications'"
    ws.append(["Status", "Today", "All earlier", "Total"])
    for i, s in enumerate(STATUSES, 2):
        ws.append([s, f'=COUNTIF({today_ref}!{t_col}:{t_col},A{i})', f'=COUNTIF({all_ref}!{a_col}:{a_col},A{i})',
                   f"=B{i}+C{i}"])
    n = len(STATUSES) + 2
    ws.append(["Total shortlisted", f"=SUM(B2:B{n - 1})", f"=SUM(C2:C{n - 1})", f"=B{n}+C{n}"])
    ws.append([])
    ws.append(["Follow-ups due (status Applied, applied 7+ days ago)",
               f'=COUNTIFS({all_ref}!{a_col}:{a_col},"Applied",{all_ref}!{f_col}:{f_col},"<="&TODAY())'])
    for c in ws[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), HEAD_FILL
    for cell in (ws.cell(n, 1), ws.cell(n + 2, 1)):
        cell.font = Font(bold=True)
    ws.column_dimensions["A"].width = 50
    for col in "BCD":
        ws.column_dimensions[col].width = 12


def _save(wb):
    paths.OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    if paths.TRACKER_FILE.exists():
        paths.BACKUP_DIR.mkdir(parents=True, exist_ok=True)
        shutil.copy2(paths.TRACKER_FILE, paths.BACKUP_DIR / f"Job Tracker {dt.datetime.now():%Y-%m-%d %H%M%S}.xlsx")
        for old in sorted(paths.BACKUP_DIR.glob("Job Tracker *.xlsx"))[:-KEEP_BACKUPS]:
            old.unlink()
    tmp = paths.TRACKER_FILE.with_name("~build " + paths.TRACKER_FILE.name)
    wb.save(tmp)
    try:
        os.replace(tmp, paths.TRACKER_FILE)
    except PermissionError:
        tmp.unlink(missing_ok=True)
        raise TrackerLocked("The tracker is open in Excel. Close it and run the build again.")


def check_not_locked():
    """Raise TrackerLocked while Excel has the workbook open (Windows refuses write access)."""
    if not paths.TRACKER_FILE.exists():
        return
    try:
        with open(paths.TRACKER_FILE, "r+b"):
            pass
    except PermissionError:
        raise TrackerLocked("The tracker is open in Excel. Close it and run the build again.")


# ---------------------------------------------------------------- build
def build(date, picks, letter_links):
    """picks: validated scored entries merged with raw jobs; letter_links: {key: path relative to output/}."""
    check_not_locked()
    if paths.TRACKER_FILE.exists():
        wb = load_workbook(paths.TRACKER_FILE)
    else:
        wb = Workbook()
        wb.remove(wb.active)
    ws_all = wb[ALL] if ALL in wb.sheetnames else _new_sheet(wb, ALL, ALL_COLS)

    carry, moved_to_all = {}, 0
    if TODAY in wb.sheetnames:
        in_all = {r.get("Key") for r in rows(ws_all) if r.get("Key")}
        for row in rows(wb[TODAY]):
            if _iso(value(row.get("Date found"))) == date:
                carry[row.get("Key")] = {f: value(row.get(f)) for f in USER_FIELDS}
            elif row.get("Key") not in in_all or not row.get("Key"):
                _append_all(ws_all, row)
                moved_to_all += 1
        wb.remove(wb[TODAY])
    _status_rules(ws_all)
    _fix_moved_letters(ws_all)

    ws = _new_sheet(wb, TODAY, TODAY_COLS, index=0)
    cv_names = config.cv_options()
    for p in picks:
        kept = carry.get(p["key"], {})
        values = {
            "Date found": dt.date.fromisoformat(date), "Fit": p["fit_score"], "Job title": p["title"],
            "Company": p["company"], "Direct / Agency": p["source_type"], "Location": p["location"],
            "Why apply": p["why"], "CV to send": os.path.basename(cv_names[p["cv"]]), "Expected salary": p["salary"],
            "Meets target?": p["meets_target"], "Visa / right to work": p["visa_notes"],
            "Deadline": p.get("deadline", ""), "Status": kept.get("Status") or "To apply",
            "Date applied": as_date(kept.get("Date applied")), "Notes": kept.get("Notes"), "Key": p["key"],
        }
        links = {"Advert": (p["link"], "View advert"), "CV to send": (_cv_link(p["cv"]), values["CV to send"])}
        if _employer_link(p):
            links["Employer site"] = (_employer_link(p), "Employer site")
        if p["key"] in letter_links:
            links["Cover letter"] = (_file_link(letter_links[p["key"]]), "Open letter")
        _put(ws, values, links, dates=("Date found", "Date applied"), wrap=True)
    wb.active = 0
    _summary(wb)
    _save(wb)
    return moved_to_all
