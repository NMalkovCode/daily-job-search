"""What each CLI command does (see __main__.py for the list)."""
import datetime as dt
import importlib
import sys

from . import config, fetch, letters, paths, scored, tracker, triage
from .filters import Filters
from .sources import SOURCES, missing_secrets
from .store import find_jobs, load_day, raw_dates


def run(args):
    date = getattr(args, "date", None) or paths.today()
    try:
        if args.command == "fetch":
            fetch.run(again=args.again, only=set(args.only.split(",")) if args.only else None)
        elif args.command == "triage":
            triage.write(date)
        elif args.command == "daily":
            fetch.run()
            triage.write(paths.today())
        elif args.command == "show":
            show(args.keys, args.date, args.chars)
        elif args.command == "validate":
            return 0 if scored.report(date) else 1
        elif args.command == "build":
            return build(date)
        elif args.command == "status":
            status()
        elif args.command == "doctor":
            return doctor()
    except tracker.TrackerLocked as e:
        print(f"ERROR: {e}")
        return 2
    except FileNotFoundError as e:
        print(f"ERROR: {e}")
        return 2
    return 0


# ---------------------------------------------------------------- show
def show(keys, date, chars):
    settings = config.settings()
    filters = Filters(settings)
    found = find_jobs(keys, date)
    for key in keys:
        j = found.get(key)
        print("=" * 100)
        if j is None:
            print(f"{key}: not found in data/raw (last 7 days)")
            continue
        details = ", ".join(f"{k}: {v}" for k, v in (j.get("details") or {}).items())
        print(f"{j['key']} | {j['title']} | {j['company']} | {j['location']} ({j['region']})")
        print(f"found {j['found_on']} | posted {j.get('posted') or '?'} | "
              f"{'AGENCY' if filters.is_agency(j) else 'DIRECT'} | flags: {', '.join(filters.flags(j)) or '-'}"
              + (f" | {details}" if details else ""))
        if j.get("salary"):
            print(f"salary (from source): {j['salary']}")
        print(f"advert: {j['link']}")
        for a in j.get("also_on", []):
            print(f"also on {a['source']}: {a['link']}")
        reason = filters.reject_reason(j)
        if reason:
            print(f"NOTE: fails hard filter -> {reason}")
        text = j.get("description") or "(no description)"
        print("-" * 100)
        print(text if len(text) <= chars else text[:chars] + f"\n[... {len(text) - chars} more characters]")


# ---------------------------------------------------------------- build
def build(date):
    errors, warnings, picks = scored.validate(date)
    for w in warnings:
        print(f"WARNING  {w}")
    if errors:
        for e in errors:
            print(f"ERROR    {e}")
        print(f"Build stopped: fix the {len(errors)} errors in data/scored/{date}.json and run build again.")
        return 1
    tracker.check_not_locked()  # before anything moves, so a locked file leaves no half-done build
    letters.rotate(date)
    links = letters.save_all(date, picks)
    moved = tracker.build(date, picks, links)
    print(f"Tracker updated: {len(picks)} jobs for {date}"
          + (f", {moved} earlier rows moved to 'All applications'" if moved else "")
          + f".\n  {paths.TRACKER_FILE}\n  {paths.LETTERS_DIR / date}")
    return 0


# ---------------------------------------------------------------- status
def status():
    today = paths.today()
    print(f"Today: {today}")
    dates = raw_dates()
    if dates:
        last = load_day(dates[-1])
        run = last["runs"][-1] if last["runs"] else {}
        state = "finished" if run.get("finished") else "NOT finished (run fetch again to resume)"
        print(f"Last fetch: {dates[-1]}, {len(last['jobs'])} jobs, {state}")
        for name, s in run.get("sources", {}).items():
            print(f"  {name:16} " + (s["skipped"] if "skipped" in s else
                                     f"{s['new']} new" + (f"  ERROR {s['error']}" if s.get("error") else "")))
    else:
        print("No fetch yet.")
    for label, path in (("Triage", paths.triage_file(today)), ("AI picks", paths.scored_file(today))):
        print(f"{label:10} {'done' if path.exists() else 'not yet'}  ({path.relative_to(paths.ROOT)})")
    if not paths.TRACKER_FILE.exists():
        print("Tracker: not created yet")
        return
    from openpyxl import load_workbook
    wb = load_workbook(paths.TRACKER_FILE, data_only=True)
    counts, due = {}, []
    for name in (tracker.TODAY, tracker.ALL):
        if name not in wb.sheetnames:
            continue
        for row in tracker.rows(wb[name]):
            st = tracker.value(row.get("Status")) or "To apply"
            counts[st] = counts.get(st, 0) + 1
            applied = tracker.as_date(tracker.value(row.get("Date applied")))
            if st == "Applied" and isinstance(applied, dt.date) and \
                    (dt.date.today() - applied).days >= tracker.FOLLOW_UP_DAYS:
                due.append(f"{tracker.value(row.get('Job title'))} - {tracker.value(row.get('Company'))} "
                           f"(applied {applied})")
    print("Tracker: " + ", ".join(f"{k} {v}" for k, v in sorted(counts.items())))
    if due:
        print("Follow-ups due:")
        for d in due:
            print(f"  {d}")


# ---------------------------------------------------------------- doctor
def doctor():
    ok = True

    def check(good, text, hint=""):
        nonlocal ok
        ok &= bool(good)
        print(f"{'OK  ' if good else 'FAIL'}  {text}" + (f"\n      -> {hint}" if hint and not good else ""))

    check(sys.version_info >= (3, 10), f"Python {sys.version.split()[0]}", "install Python 3.10 or newer")
    for module in ("requests", "bs4", "openpyxl", "docx"):
        try:
            importlib.import_module(module)
            check(True, f"package {module}")
        except ImportError:
            check(False, f"package {module}", "run setup (setup.bat or ./setup.sh)")
    try:
        settings = config.settings()
        Filters(settings)
        check(True, "config/settings.json")
    except Exception as e:
        check(False, f"config/settings.json: {e}", "run setup, or fix the JSON (a missing comma or quote?)")
        return 1
    if settings["candidate"].get("name", "Your Name") == "Your Name":  # still the example's placeholder
        print("WARN  your name and signature are still the example's: edit candidate in config/settings.json")
    check(paths.PROFILE_FILE.exists(), "config/profile.md", "run setup (it copies config/profile.example.md)")
    if paths.PROFILE_FILE.exists() and "<Your name>" in paths.PROFILE_FILE.read_text(encoding="utf-8"):
        print("WARN  config/profile.md is still the template: fill it in (Claude can help, see README.md)")
    print(f"--    countries searched: {', '.join(settings['regions'])}")
    try:
        from .sources import ats
        check(True, f"config/companies.json ({len(ats.companies())} employer boards)")
    except Exception as e:
        check(False, f"config/companies.json: {e}")
    for name, module in SOURCES.items():
        enabled = settings["sources"].get(name, {}).get("enabled", False)
        missing = missing_secrets(module, config.secrets())
        if not enabled:
            print(f"--    source {name}: disabled in settings.json")
        elif missing:
            print(f"--    source {name}: no API key, will be skipped (add {', '.join(missing)} to config/secrets.json)")
        else:
            check(True, f"source {name}")
    for cv_id, rel in config.cv_options().items():
        if config.cv_path(cv_id).exists():
            check(True, f"CV '{cv_id}': {rel}")
        else:  # not fatal: the tracker just cannot link the file
            print(f"WARN  CV '{cv_id}' not found: {config.cv_path(cv_id)}\n"
                  f"      -> copy it there, or change candidate.cvs / cv_root in config/settings.json")
    try:
        tracker.check_not_locked()
        check(True, "tracker is not open in Excel")
    except tracker.TrackerLocked as e:
        check(False, str(e))
    print("All good." if ok else "Some checks failed - see the hints above.")
    return 0 if ok else 1
