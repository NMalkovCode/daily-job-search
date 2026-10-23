"""validate + build end to end: tracker rows, carried statuses, letters, rotation."""
import datetime as dt

from openpyxl import load_workbook

from conftest import job, pick, write_day, write_scored
from jobsearch import commands, letters, paths, scored, tracker

D1, D2 = "2026-10-20", "2026-10-21"


def sheet_rows(name):
    wb = load_workbook(paths.TRACKER_FILE)
    return tracker.rows(wb[name])


def set_status(key, status, applied=None):
    wb = load_workbook(paths.TRACKER_FILE)
    ws = wb[tracker.TODAY]
    heads = tracker._headers(ws)
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, heads["Key"]).value == key:
            ws.cell(r, heads["Status"], status)
            if applied:
                ws.cell(r, heads["Date applied"], applied)
    wb.save(paths.TRACKER_FILE)


def test_validation_errors(project):
    write_day(D1, [job(1), job(2, title="Pricing Analyst")])
    write_scored(D1, [
        pick("li:1"), pick("li:1"),                          # duplicate
        pick("li:404"),                                       # unknown key
        pick("li:2", cv="nonexistent", cover_letter="Dear [Hiring Manager], ..."),
        {"key": "li:2", "recommend": False, "fit_score": 3, "why": "Too senior."},
    ])
    errors, warnings, picks = scored.validate(D1)
    text = "\n".join(errors)
    assert "twice" in text and "not found" in text and "cv 'nonexistent'" in text and "placeholder" in text
    assert [p["key"] for p in picks] == ["li:1", "li:1", "li:2"]


def test_sign_off_is_stripped():
    assert scored.strip_sign_off("Dear X,\n\nBody.\n\nKind regards,\nAlex") == "Dear X,\n\nBody."


def test_build_then_next_day_moves_rows_with_status(project):
    write_day(D1, [job(1), job(2, title="Pricing Analyst", company="Beta")])
    write_scored(D1, [pick("li:1", fit_score=9), pick("li:2", company="Beta Real Employer")])
    assert commands.build(D1) == 0

    today = sheet_rows(tracker.TODAY)
    assert [r["Key"] for r in today] == ["li:1", "li:2"]
    assert today[1]["Company"] == "Beta Real Employer"          # AI's employer name wins
    assert today[0]["Advert"][1] == "https://example.com/jobs/1"
    letter = today[0]["Cover letter"][1]
    assert (paths.OUTPUT_DIR / letter.replace("%20", " ")).exists()

    # same-day rebuild keeps what the user typed
    set_status("li:1", "Applied", dt.datetime(2026, 10, 20))
    assert commands.build(D1) == 0
    assert sheet_rows(tracker.TODAY)[0]["Status"] == "Applied"

    # next day: yesterday's rows move to All applications, status and date kept
    write_day(D2, [job(3, title="Strategy Analyst", company="Gamma")])
    write_scored(D2, [pick("li:3", company="Gamma")])
    assert commands.build(D2) == 0
    all_rows = sheet_rows(tracker.ALL)
    assert [r["Key"] for r in all_rows] == ["li:1", "li:2"]
    assert all_rows[0]["Status"] == "Applied" and all_rows[1]["Status"] == "To apply"
    assert tracker.as_date(all_rows[0]["Date applied"]) == dt.date(2026, 10, 20)
    assert str(all_rows[0]["Follow up by"]).startswith("=IF(")
    assert [r["Key"] for r in sheet_rows(tracker.TODAY)] == ["li:3"]

    # a job already in the tracker cannot be recommended again (same key or a repost)
    write_day("2026-10-22", [job(4, title="Strategy Analyst", company="Gamma")])
    write_scored("2026-10-22", [pick("li:4", company="Gamma")])
    errors, _, _ = scored.validate("2026-10-22")
    assert any("already in the tracker" in e for e in errors)

    # building twice never duplicates the status drop-down
    wb = load_workbook(paths.TRACKER_FILE)
    assert len(wb[tracker.ALL].data_validations.dataValidation) == 1


def test_edited_letter_is_kept(project):
    write_day(D1, [job(1)])
    write_scored(D1, [pick("li:1")])
    commands.build(D1)
    path = next((paths.LETTERS_DIR / D1).glob("*.docx"))
    path.write_bytes(path.read_bytes() + b"edited")
    write_scored(D1, [pick("li:1", cover_letter=pick("li:1")["cover_letter"].replace("Dear", "Hello"))])
    commands.build(D1)
    assert path.read_bytes().endswith(b"edited")


def test_rotation_moves_all_files_and_fixes_links(project):
    write_day(D1, [job(1)])
    write_scored(D1, [pick("li:1")])
    commands.build(D1)
    folder = paths.LETTERS_DIR / D1
    (folder / "my letter.pdf").write_bytes(b"%PDF")

    later = "2026-10-28"
    write_day(later, [job(5, title="Data Analyst", company="Delta")])
    write_scored(later, [pick("li:5", company="Delta")])
    assert commands.build(later) == 0
    assert not folder.exists()
    prev = paths.PREVIOUS_LETTERS_DIR
    assert (prev / f"{D1} - my letter.pdf").exists()
    moved_row = sheet_rows(tracker.ALL)[0]
    target = moved_row["Cover letter"][1].replace("%20", " ")
    assert target.startswith("Cover Letters/Previous Cover Letters/") and (paths.OUTPUT_DIR / target).exists()


def test_locked_tracker_stops_before_anything_moves(project, monkeypatch):
    write_day(D1, [job(1)])
    write_scored(D1, [pick("li:1")])
    commands.build(D1)

    def locked():
        raise tracker.TrackerLocked("open in Excel")
    monkeypatch.setattr(tracker, "check_not_locked", locked)
    rotated = []
    monkeypatch.setattr(letters, "rotate", lambda d: rotated.append(d))
    assert commands.run(type("A", (), {"command": "build", "date": D1})()) == 2
    assert rotated == []
